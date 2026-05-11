"""通用数据导入导出 — 基于 openpyxl"""

import logging
from io import BytesIO

from django.apps import apps
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAdminUser
from rest_framework.response import Response

logger = logging.getLogger("djangoadminx.data_center")


def _get_model(model_label):
    try:
        app_label, model_name = model_label.rsplit(".", 1)
        return apps.get_model(app_label, model_name)
    except (LookupError, ValueError):
        return None


def _get_exportable_models():
    return [
        "accounts.User", "accounts.Role",
        "menu.Menu",
        "config_center.Config",
        "cluster.ClusterNode",
        "webservice.WebService", "webservice.ScheduleJob",
    ]


@api_view(["GET"])
@permission_classes([IsAdminUser])
def export_data(request):
    model_label = request.query_params.get("model", "")
    fields_param = request.query_params.get("fields", "")

    model = _get_model(model_label)
    if model is None:
        return Response({"code": 400, "msg": "模型不存在: " + model_label})
    if model_label not in _get_exportable_models():
        return Response({"code": 403, "msg": "模型不支持导出: " + model_label})

    if fields_param:
        fields = [f.strip() for f in fields_param.split(",") if f.strip()]
    else:
        fields = [f.name for f in model._meta.fields if f.name != "id"]
        fields = [f for f in fields if f not in ("password", "encrypted_value")]
    if "id" not in fields:
        fields.insert(0, "id")

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return Response({"code": 500, "msg": "openpyxl 未安装，请 pip install openpyxl"})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = model._meta.verbose_name

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, field_name in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col, value=field_name)
        cell.fill = header_fill
        cell.font = header_font

    qs = model.objects.all()
    for row_idx, obj in enumerate(qs, 2):
        for col_idx, field_name in enumerate(fields, 1):
            val = getattr(obj, field_name, "")
            if hasattr(val, "isoformat"):
                val = val.isoformat() if val else ""
            elif hasattr(val, "pk"):
                val = str(val)
            ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = model_label.replace(".", "_") + "_export.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="' + filename + '"'
    return response


@api_view(["POST"])
@permission_classes([IsAdminUser])
def import_data(request):
    model_label = request.data.get("model", "")
    file_obj = request.FILES.get("file")

    if not model_label:
        return Response({"code": 400, "msg": "请指定 model"})
    if not file_obj:
        return Response({"code": 400, "msg": "请上传 Excel 文件"})

    model = _get_model(model_label)
    if model is None:
        return Response({"code": 400, "msg": "模型不存在: " + model_label})
    if model_label not in _get_exportable_models():
        return Response({"code": 403, "msg": "模型不支持导入: " + model_label})

    try:
        import openpyxl
    except ImportError:
        return Response({"code": 500, "msg": "openpyxl 未安装"})

    try:
        wb = openpyxl.load_workbook(file_obj)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            return Response({"code": 400, "msg": "Excel 为空或只有表头"})

        headers = [str(h) for h in rows[0]]
        success, failed = 0, 0
        errors = []

        for row_idx, row in enumerate(rows[1:], 2):
            try:
                data = {}
                for col_idx, val in enumerate(row):
                    if col_idx < len(headers):
                        data[headers[col_idx]] = val
                obj_id = data.pop("id", None)
                if obj_id:
                    model.objects.update_or_create(id=obj_id, defaults=data)
                else:
                    model.objects.create(**data)
                success += 1
            except Exception as e:
                failed += 1
                errors.append("第%d行: %s" % (row_idx, str(e)[:200]))

        return Response({
            "code": 200,
            "msg": "success",
            "data": {
                "success": success,
                "failed": failed,
                "errors": errors[:20],
            },
        })
    except Exception as e:
        err_msg = str(e)[:500]
        return Response({"code": 500, "msg": "导入失败: " + err_msg})